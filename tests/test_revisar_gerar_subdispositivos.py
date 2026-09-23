"""Botão 'Revisar e gerar DOCX' (atalho análise→correção) e preservação de
subdispositivos (§/incisos) ao alterar apenas o caput de um parágrafo físico do
Word que contém caput + § no MESMO elemento ``<w:p>``.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from sejus_project.agent import agent
from sejus_project.tools.document_infra import docx_builder, modelos
from sejus_project.tools.llm_tools import analysis_registry
from sejus_project.tools.llm_tools import document_generation as generation
from sejus_project.tools.llm_tools import document_improvement as minuta
from sejus_project.tools.llm_tools import user_files

_CAPUT = (
    "Art. 10 O recebimento e a entrega de materiais de artesanato fornecidos "
    "por familiares ou terceiros na unidade penal deverão obedecer estritamente "
    "às normas técnicas e de segurança da portaria de visitas vigente."
)
_CAPUT_NOVO = _CAPUT.replace("unidade penal", "Unidade Penal")
_PAR1 = "§ 1º É obrigatória a apresentação de carteira de visitante válida."
_PAR2 = "§ 2º Não será permitido o recebimento de materiais de visitante vencido."


def _docx_caput_com_paragrafos(path: Path) -> Path:
    """DOCX com caput + §§ no MESMO parágrafo físico (separados por <w:br/>)."""
    doc = Document()
    doc.add_paragraph("INSTRUÇÃO NORMATIVA Nº 1/2026")
    p = doc.add_paragraph()
    p.add_run(_CAPUT)
    r = p.add_run()
    r.add_break()
    r.add_text(_PAR1)
    r2 = p.add_run()
    r2.add_break()
    r2.add_text(_PAR2)
    doc.save(str(path))
    return path


def _runs(w_p):
    return list(w_p.findall(qn("w:r")))


def _tachado(run) -> bool:
    rpr = run.find(qn("w:rPr"))
    strike = rpr.find(qn("w:strike")) if rpr is not None else None
    return strike is not None and (strike.get(qn("w:val")) or "") not in ("0", "false")


def _verde(run) -> bool:
    rpr = run.find(qn("w:rPr"))
    cor = rpr.find(qn("w:color")) if rpr is not None else None
    return cor is not None and (cor.get(qn("w:val")) or "").upper() == "2E7D32"


def test_alterar_caput_preserva_paragrafos_no_mesmo_wp(tmp_path):
    origem = _docx_caput_com_paragrafos(tmp_path / "doc.docx")
    texto = "\n".join(p.text for p in Document(str(origem)).paragraphs)
    perfil = modelos.crear_perfil_de_arquivo(origem, texto, "doc")

    alteracao = {
        "tipo": "alterado",
        "rotulo": "Art. 10",
        "trecho_original": _CAPUT,
        "novo_texto": _CAPUT_NOVO,
    }
    saida = docx_builder.montar_docx_revisado(perfil, [alteracao], [], [], tmp_path)

    doc = Document(str(saida))
    alvo = next(
        wp for wp in (ch for ch in doc.element.body if ch.tag == qn("w:p"))
        if _PAR1 in "".join((t.text or "") for t in wp.iter(qn("w:t")))
    )
    ativos = [r for r in _runs(alvo) if not _tachado(r)]
    texto_ativo = "".join(
        "".join((t.text or "") for t in r.findall(qn("w:t"))) for r in ativos
    )
    # Novo caput ativo e os dois §§ preservados (não tachados).
    assert "Unidade Penal" in texto_ativo
    assert _PAR1 in texto_ativo
    assert _PAR2 in texto_ativo
    # O caput antigo ficou tachado (marcação de substituição).
    antigos = [r for r in _runs(alvo) if _tachado(r)]
    assert any("unidade penal" in "".join((t.text or "") for t in r.findall(qn("w:t"))) for r in antigos)
    # Nenhum § foi tachado.
    for r in _runs(alvo):
        txt = "".join((t.text or "") for t in r.findall(qn("w:t")))
        if "§" in txt:
            assert not _tachado(r)


def test_perda_de_subitens_detectada_e_alteracao_descartada(tmp_path):
    conteudo = "\n".join([_CAPUT, _PAR1, _PAR2])
    com_perda = {
        "tipo": "alterado",
        "rotulo": "Art. 10",
        "trecho_original": f"{_CAPUT}\n{_PAR1}\n{_PAR2}",
        "novo_texto": _CAPUT_NOVO,
    }
    explicita = dict(com_perda, subitens_removidos=["§ 1º", "§ 2º"])
    valida = {
        "tipo": "alterado",
        "rotulo": "Art. 11",
        "trecho_original": "Art. 11 Texto original.",
        "novo_texto": "Art. 11 Texto revisado.",
    }
    assert minuta._subitens_perdidos(com_perda) == ["§ 1º", "§ 2º"]
    assert minuta._subitens_perdidos(explicita) == []
    assert minuta._subitens_perdidos(valida) == []

    conteudo_ok = conteudo + "\nArt. 11 Texto original."
    alt, _, _, descartados = minuta._filtrar_patch_valido(
        conteudo_ok, [com_perda, valida], [], []
    )
    rotulos = [a["rotulo"] for a in alt]
    assert "Art. 11" in rotulos and "Art. 10" not in rotulos
    assert any("perda involuntaria" in d for d in descartados)


def _preparar_pasta(monkeypatch, tmp_path):
    monkeypatch.setattr(user_files, "IMPORTACOES_DIR", tmp_path)
    return tmp_path


def _registrar_analise_existente(nome: str, apontamentos: list[dict]) -> str:
    """Registra análise para o arquivo (mesma sessão + versão) e devolve o nome."""
    sessao = analysis_registry.sessao_atual()
    caminho = user_files._resolve_file(nome)
    texto = user_files.extract_file_text(caminho)
    analysis_registry.registrar(
        sessao, caminho.name, analysis_registry.hash_conteudo(texto),
        "análise registrada", apontamentos,
    )
    return nome


def test_revisar_gerar_reutiliza_analise_da_mesma_versao(monkeypatch, tmp_path):
    pasta = _preparar_pasta(monkeypatch, tmp_path)
    (pasta / "doc.txt").write_text("Art. 1º Texto.", encoding="utf-8")
    analysis_registry.nova_sessao()
    _registrar_analise_existente("doc.txt", [{"id": "ap-1", "texto": "corrigir X"}])

    capturado = {}
    monkeypatch.setattr(
        agent, "_executar_melhoria",
        lambda **kw: capturado.update(kw) or "melhoria-ok",
    )
    monkeypatch.setattr(
        agent, "_analise_isolada",
        lambda nome: (_ for _ in ()).throw(AssertionError("não deveria analisar")),
    )

    resposta = agent._tratar_revisao_direta("Revise e gere o DOCX do arquivo 'doc.txt'")
    assert resposta == "melhoria-ok"
    assert capturado["filename"] == "doc.txt"
    assert [a["id"] for a in capturado["apontamentos"]] == ["ap-1"]


def test_revisar_gerar_analisa_quando_nao_ha_analise(monkeypatch, tmp_path):
    pasta = _preparar_pasta(monkeypatch, tmp_path)
    (pasta / "novo.txt").write_text("Art. 1º Texto.", encoding="utf-8")
    analysis_registry.nova_sessao()

    def _analise_isolada(nome):
        _registrar_analise_existente(nome, [{"id": "ap-9", "texto": "faltava prazo"}])

    capturado = {}
    monkeypatch.setattr(agent, "_analise_isolada", _analise_isolada)
    monkeypatch.setattr(
        agent, "_executar_melhoria",
        lambda **kw: capturado.update(kw) or "melhoria-ok",
    )

    resposta = agent._tratar_revisao_direta("Revise e gere o DOCX do arquivo 'novo.txt'")
    assert resposta == "melhoria-ok"
    assert [a["id"] for a in capturado["apontamentos"]] == ["ap-9"]


def test_analise_antiga_nao_reutilizada_quando_arquivo_muda(monkeypatch, tmp_path):
    pasta = _preparar_pasta(monkeypatch, tmp_path)
    arquivo = pasta / "doc.txt"
    arquivo.write_text("Art. 1º Versão antiga.", encoding="utf-8")
    analysis_registry.nova_sessao()
    _registrar_analise_existente("doc.txt", [{"id": "ap-1", "texto": "x"}])

    # Arquivo modificado com o MESMO nome: hash diferente → análise não vale.
    arquivo.write_text("Art. 1º Versão NOVA.", encoding="utf-8")
    assert generation.analise_para_correcao("doc.txt") is None


def test_chat_e_botao_compartilham_mesma_analise_e_apontamentos(monkeypatch, tmp_path):
    pasta = _preparar_pasta(monkeypatch, tmp_path)
    (pasta / "doc.txt").write_text("Art. 1º Texto.", encoding="utf-8")
    analysis_registry.nova_sessao()
    _registrar_analise_existente(
        "doc.txt", [{"id": "ap-1", "texto": "corrigir X"}]
    )

    capturas: list[dict] = []
    monkeypatch.setattr(
        agent, "_executar_melhoria",
        lambda **kw: capturas.append(dict(kw)) or "melhoria-ok",
    )

    # Caminho do BOTÃO.
    agent._tratar_revisao_direta("Revise e gere o DOCX do arquivo 'doc.txt'")
    # Caminho do CHAT (confirmação curta após a análise).
    agent._tratar_correcao_direta("sim")

    assert len(capturas) == 2
    assert capturas[0]["filename"] == capturas[1]["filename"] == "doc.txt"
    assert [a["id"] for a in capturas[0]["apontamentos"]] == [
        a["id"] for a in capturas[1]["apontamentos"]
    ] == ["ap-1"]

