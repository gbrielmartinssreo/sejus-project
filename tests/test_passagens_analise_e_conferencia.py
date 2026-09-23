"""Os 4 itens do feedback sobre a melhoria de documentos:

1. Registro em separado das 5 passagens de ANÁLISE (com chamadas/resultado/
   duração), distintas das passagens pós-geração;
2. Conferência de cada achado contra o ORIGINAL (trecho + localização),
   descartando o que o contradiz (incluindo 'parafins'/'parágrafo único do 20'
   e elogios) antes da melhoria;
3. Reconciliar a cobertura com o DOCX final: correções automáticas na mesma
   rastreabilidade, status pelo resultado efetivo (capítulos I..N não podem
   terminar 'nao_aplicado'), sem elogios na lista de tarefas;
4. Preservar o objetivo do achado e anotar comentários com achado_id +
   justificativa.
"""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

from sejus_project.tools.document_infra import analise_validacao as av
from sejus_project.tools.document_infra.docx_builder import (
    _aplicar_automatica_inline,
    _texto_comentario,
    montar_docx_revisado,
)
from sejus_project.tools.document_infra.docx_engine import paragraph_text
from sejus_project.tools.document_infra.modelos import PORTARIA, PerfilModelo
from sejus_project.tools.llm_tools import document_generation as generation
from sejus_project.tools.llm_tools import document_improvement as di

_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

_PASSOS_ANALISE = [
    "ortografia", "estrutura", "clareza", "fundamentação",
    "conferência dos achados",
]


def _conteudo_base(extra: str = "") -> str:
    return (
        "INSTRUÇÃO NORMATIVA Nº 1/2026\n"
        "Art. 20 Este texto orienta o trabalho artesanal no âmbito da SEJUS.\n"
        "Art. 21 Esta Instrução Normativa entra em vigor na data da publicação."
        + (f"\n{extra}" if extra else "")
    )


def _perfil_teste(tmp_path: Path, conteudo: str) -> PerfilModelo:
    from docx import Document

    modelo = tmp_path / "USUARIO_ARTESANATO.docx"
    doc = Document()
    for linha in conteudo.splitlines():
        doc.add_paragraph(linha)
    doc.save(str(modelo))
    return PerfilModelo(
        name="USUARIO_Teste", file=str(modelo),
        act_types=PORTARIA.act_types, patterns=PORTARIA.patterns,
        preservar_moldura=True,
    )


# ---------------------------------------------------------------------------
# 1. Passagens de análise registradas em separado
# ---------------------------------------------------------------------------


def test_rodar_passagens_analise_registra_5_passos_com_chamadas_e_duracao():
    res = av.rodar_passagens_analise(_conteudo_base(), [])
    assert [p["passo"] for p in res["passagens"]] == _PASSOS_ANALISE
    for p in res["passagens"]:
        assert isinstance(p.get("tempo_ms"), (int, float))
        assert p.get("resultado")
        assert p.get("chamadas"), f"{p['passo']} sem chamadas"
        assert all(c.get("funcao") for c in p["chamadas"])


def test_passagens_analise_distintas_das_pos_geracao():
    """Os passos de ANÁLISE não colidem com os rótulos das pós-geração."""
    res = av.rodar_passagens_analise(_conteudo_base(), [])
    rotulos_analise = {p["passo"] for p in res["passagens"]}
    rotulos_pos = {
        "1. Considerandos preservados", "2. Sequência de capítulos",
        "3. Sem duplicações introduzidas",
        "4. Correspondência relatório x texto", "5. Art. 15 e §§ preservados",
    }
    assert rotulos_analise == set(_PASSOS_ANALISE)
    assert rotulos_analise.isdisjoint(rotulos_pos)


# ---------------------------------------------------------------------------
# 2. Conferência dos achados
# ---------------------------------------------------------------------------


def test_conferencia_descarta_contradicoes_do_feedback():
    """Casos do feedback reprodutor: ementa 'não menciona' o que ela menciona,
    inclusão de Lei já citada e repetição que não existe no trecho."""
    conteudo = (
        "INSTRUÇÃO NORMATIVA Nº 1/2026\n"
        "EMENTA: Estabelece os critérios para o fluxo financeiro e a "
        "governança do trabalho artesanal no âmbito da SEJUS.\n"
        "CONSIDERANDO a Lei Federal nº 4.320, de 17 de março de 1964;\n"
        "Art. 9º Os materiais serão fornecidos exclusivamente pela família.\n"
        "Art. 10 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    apontamentos = [
        {"id": "a-ementa", "texto": "A ementa não menciona o fluxo financeiro "
         "e a governança do trabalho artesanal."},
        {"id": "a-lei", "texto": "Recomenda-se incluir a Lei 4.320/1964 nos "
         "considerandos."},
        {"id": "a-repet", "texto": "Há repetição da palavra \"exclusivamente\" "
         "no art. 9º."},
    ]
    validos, descartados = av.conferir_achados(conteudo, apontamentos)
    assert {d["apontamento_id"] for d in descartados} == {
        "a-ementa", "a-lei", "a-repet",
    }
    for d in descartados:
        assert d["motivo"]
        assert d["localizacao"]
    assert validos == []


def test_conferencia_mantem_achados_confirmados_no_original():
    conteudo = (
        "INSTRUÇÃO NORMATIVA Nº 1/2026\n"
        "CAPÍTULO II\nDO ALCANCE\n"
        "CAPÍTULO III\nDO REGISTRO\n"
        "CAPÍTULO III\nDO CONTROLE\n"
        "Art. 15 O exercício da atividade de auditoria será registrado.\n"
        "Art. 20 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    apontamentos = [
        {"id": "ok-repet", "texto": "Repetição do título \"CAPÍTULO III\": há "
         "dois capítulos III distintos."},
        {"id": "ok-ausente", "texto": "A IN não prevê o recurso administrativo "
         "contra as decisões da unidade."},
    ]
    validos, descartados = av.conferir_achados(conteudo, apontamentos)
    ids_validos = {v["id"] for v in validos}
    # Duplicação REAL de capítulo é mantida; ausência confirmada fica adiável.
    assert ids_validos == {"ok-repet", "ok-ausente"}
    assert not any(d["apontamento_id"] == "ok-repet" for d in descartados)


def test_conferencia_descarta_elogio_de_conformidade_e_mantem_acao():
    validos, descartados = av.conferir_achados(
        _conteudo_base(),
        [
            {"id": "elogio", "texto": "Fundamentação legal abrangente e "
             "atualizada na IN."},
            {"id": "ac", "texto": "Sugiro ajustar a fundamentação do "
             "preâmbulo."},
        ],
    )
    assert {d["apontamento_id"] for d in descartados} == {"elogio"}
    # 'ac' não tem cláusula negada nem trecho a conferir → permanece acionável.
    assert [v["id"] for v in validos] == ["ac"]


def test_conferencia_descarta_falta_vigencia_quando_doc_ja_preve():
    conteudo = (
        "INSTRUÇÃO NORMATIVA Nº 1/2026\n"
        "Art. 33 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    validos, descartados = av.conferir_achados(
        conteudo,
        [{"id": "vig", "texto": "Falta incluir o artigo de vigência explícito."}],
    )
    assert validos == []
    assert descartados and "vigor" in descartados[0]["motivo"]
    # Sem dispositivo de vigor (sem o lexema 'vigor'/'vigência'), o achado
    # continua acionável.
    sem_vigor = conteudo.replace(
        "entra em vigor na data da publicação",
        "será regulamentada em até 90 dias",
    )
    validos2, _ = av.conferir_achados(
        sem_vigor,
        [{"id": "vig", "texto": "Falta incluir o artigo de vigência explícito."}],
    )
    assert [v["id"] for v in validos2] == ["vig"]


def test_conferencia_sem_nada_descartado_devolve_todos_validos():
    """Regressão: quando nada é descartado, os válidos não podem sumir."""
    conteudo = (
        "INSTRUÇÃO NORMATIVA Nº 1/2026\n"
        "Art. 19 A comercialização poderá ser realizada por:\n"
        "II - lojas físicas ou plataformas virtuais.\n"
        "Art. 20 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    apontamentos = [
        {"id": "a1", "texto": "Reformular os incisos II e IV do art. 19."},
        {"id": "a2", "texto": "Incluir prazo de validade da autorização."},
    ]
    res = av.rodar_passagens_analise(conteudo, apontamentos)
    assert {a["id"] for a in res["apontamentos_validos"]} == {"a1", "a2"}
    assert res["achados_descartados"] == []


# ---------------------------------------------------------------------------
# 3. Correções automáticas + reconciliação da cobertura
# ---------------------------------------------------------------------------


def test_detectar_correcoes_redacao_gera_itens_de_patch():
    conteudo = (
        "Art. 15 O artesanato será reconhecido parafins de remição de pena.\n"
        "§ 1º Na hipótese do caput, valem as formas previstas no parágrafo "
        "único do 20 desta Instrução Normativa."
    )
    itens = av.detectar_correcoes_redacao(conteudo)
    pares = {(i["buscar"], i["substituir"]) for i in itens}
    assert ("parafins", "para fins") in pares
    assert ("parágrafo único do 20", "parágrafo único do art. 20") in pares
    for item in itens:
        assert item["tipo"] == "corrigido"
        assert item["automatica"] is True
        assert item["trecho_original"].strip() and item["novo_texto"].strip()
        assert di._item_tem_ancora(conteudo, item)


def test_correcoes_automaticas_entram_no_patch_valido():
    conteudo = (
        "Art. 1º Esta Instrução Normativa vale para parafins de padronização.\n"
        "Art. 2º Esta Instrução Normativa entra em vigor na data da publicação."
    )
    alteracoes = di._dedupe_mudancas(av.detectar_correcoes_redacao(conteudo))
    validos, _, _, descartados = di._filtrar_patch_valido(
        conteudo, alteracoes, [], []
    )
    assert len(validos) == 1
    assert validos[0]["buscar"] == "parafins"
    assert descartados == []


def test_reconciliar_cobertura_status_pelo_resultado_efetivo():
    auto = av.detectar_correcoes_redacao(
        "Art. 1º Texto com parafins de exemplo.\n"
        "Art. 2º Esta Instrução Normativa entra em vigor na data da publicação."
    )
    renum = [{
        "tipo": "corrigido", "rotulo": "CAPÍTULO III",
        "trecho_original": "CAPÍTULO III\nDO REGISTRO",
        "novo_texto": "CAPÍTULO III\nDO REGISTRO",
        "renumeracao_engine": True,
        "detalhe": "renumeração canônica de capítulos aplicada",
    }]
    cobertura = [
        {
            "apontamento_id": "a-capitulo",
            "apontamento": "Repetição do título CAPÍTULO III; sugere-se "
            "renumerar os capítulos.",
            "origem": "analise", "status": "nao_aplicado",
            "referencia": "", "motivo": "anterior",
        },
        {
            "apontamento_id": "a-ortografia",
            "apontamento": "Corrigir erro de ortografia: parafins deve ser "
            "separado.",
            "origem": "analise", "status": "falhou",
            "referencia": "", "motivo": "não foi alterado",
        },
        {
            "apontamento_id": "a-outro",
            "apontamento": "Adicionar cláusula de vigência.",
            "origem": "analise", "status": "nao_aplicado",
            "referencia": "", "motivo": "escopo fora do ato",
        },
    ]
    out = di._reconciliar_cobertura_automatica(cobertura, auto + renum, [])
    por_id = {c["apontamento_id"]: c for c in out}
    assert por_id["a-capitulo"]["status"] == "aplicado"
    assert "correção automática do sistema" in por_id["a-capitulo"]["motivo"]
    assert por_id["a-ortografia"]["status"] == "aplicado"
    assert por_id["a-outro"]["status"] == "nao_aplicado"


def test_reconciliar_cobertura_registra_linha_auto():
    """Correção automática sem apontamento correspondente ganha rastreabilidade
    própria (na mesma lista de cobertura), com status pelo resultado efetivo."""
    auto = av.detectar_correcoes_redacao(
        "Art. 1º Segue o parágrafo único do 20 desta norma.\n"
        "Art. 20 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    assert auto
    out = di._reconciliar_cobertura_automatica([], auto, [])
    auto_rows = [c for c in out if c["apontamento_id"].startswith("auto-")]
    assert len(auto_rows) == 1
    assert auto_rows[0]["status"] == "aplicado"
    assert auto_rows[0]["origem"] == "sistema (correção automática)"
    assert auto_rows[0]["referencia"]


# ---------------------------------------------------------------------------
# 4. DOCX: renderização cirúrgica + comentário com justificativa/achado_id
# ---------------------------------------------------------------------------


def _item_correcao_automatica(base: str) -> dict:
    return {
        "tipo": "corrigido", "rotulo": base[:70],
        "trecho_original": base,
        "novo_texto": base.replace("parágrafo único do 20",
                                   "parágrafo único do art. 20"),
        "buscar": "parágrafo único do 20",
        "substituir": "parágrafo único do art. 20",
        "localizacao": "parágrafo que remete ao art. 20",
        "detalhe": "referência: 'parágrafo único do 20' deve indicar o "
        "dispositivo ('art. 20').",
        "automatica": True, "origem": "sistema (correção automática)",
    }


def test_docx_renderiza_correcao_automatica_inline(tmp_path):
    from docx import Document

    base = "Art. 20 Todos os valores seguem o parágrafo único do 20 desta norma."
    conteudo = base + "\n" + (
        "Art. 21 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    perfil = _perfil_teste(tmp_path, conteudo)
    saida = montar_docx_revisado(perfil, [_item_correcao_automatica(base)], [], [],
                                 tmp_path / "saida")
    doc_final = Document(str(saida))
    alvo = next(
        p for p in doc_final.paragraphs
        if "art. 20" in p.text and "único" in p.text
    )
    correcoes = [
        "".join(t.text or "" for t in r.iter(_NS + "t"))
        for r in alvo._element.iter(_NS + "r")
    ]
    assert any(t == "parágrafo único do 20" for t in correcoes)
    assert any(t == "parágrafo único do art. 20" for t in correcoes)
    # A correção é INLINE: não nasceu um parágrafo novo/duplicado.
    alvos = [
        p for p in doc_final.paragraphs
        if "art. 20" in p.text and "único" in p.text and p.text.strip()
    ]
    assert len(alvos) == 1


def test_docx_comentario_correcao_automatica_carrega_detalhe_e_achado(tmp_path):
    import zipfile

    from docx import Document

    base = "Art. 20 Segue o parágrafo único do 20 desta norma."
    conteudo = base + "\n" + (
        "Art. 21 Esta Instrução Normativa entra em vigor na data da publicação."
    )
    perfil = _perfil_teste(tmp_path, conteudo)
    item = _item_correcao_automatica(base)
    item["achado_id"] = "ap-42"
    saida = montar_docx_revisado(perfil, [item], [], [], tmp_path / "saida")
    with zipfile.ZipFile(str(saida)) as arquivo:
        assert "word/comments.xml" in arquivo.namelist()
        xml = arquivo.read("word/comments.xml").decode("utf-8")
    assert "Correção automática do sistema" in xml
    assert "ap-42" in xml


def test_texto_comentario_automatica_inclui_justificativa():
    comentario = _texto_comentario({
        "automatica": True,
        "detalhe": "referência: 'parágrafo único do 20' corrigida.",
        "achado_id": "ap-7",
    })
    assert "Correção automática do sistema" in comentario
    assert "ap-7" in comentario


def test_sistema_melhoria_preserva_objetivo_do_achado():
    instrucoes = di._sistema_melhoria()
    assert "PRESERVE O OBJETIVO DO ACHADO" in instrucoes
    assert "nao_aplicado" in instrucoes


def test_aplicar_automatica_inline_nao_expande_paragrafo(tmp_path):
    from docx import Document

    modelo = tmp_path / "m.docx"
    doc = Document()
    doc.add_paragraph("Art. 20 Segue o parágrafo único do 20 desta norma.")
    doc.save(str(modelo))
    p = Document(str(modelo)).paragraphs[0]
    _aplicar_automatica_inline(
        p._element,
        {"buscar": "parágrafo único do 20",
         "substituir": "parágrafo único do art. 20"},
    )
    texto = paragraph_text(p._element)
    assert "parágrafo único do 20" in texto
    assert "parágrafo único do art. 20" in texto
    # antes + buscado(tachado) + substituo(verde) + restante, sem duplicar.
    assert len(p._element.findall(_NS + "r")) == 4


# ---------------------------------------------------------------------------
# Fluxo completo: melhoria registra passagens de análise e aplica a correção
# ---------------------------------------------------------------------------


def test_fluxo_melhoria_registra_passagens_analise_e_aplica_correcao_automatica(
    monkeypatch, tmp_path
):
    """Pedido de correção: as 5 passagens de análise saem em separado, o achado
    contraditório é descartado, e uma correção automática de redação entra na
    mesma rastreabilidade com status pelo resultado efetivo."""
    from docx import Document

    conteudo = _conteudo_base(
        "Art. 20 Todos os valores seguem o parágrafo único do 20 desta norma."
    )
    modelo = tmp_path / "USUARIO_ARTESANATO.docx"
    doc = Document()
    for linha in conteudo.splitlines():
        doc.add_paragraph(linha)
    doc.save(str(modelo))

    monkeypatch.setattr(generation, "_resolve_file", lambda n: modelo)
    monkeypatch.setattr(generation, "extract_file_text", lambda p: conteudo)
    monkeypatch.setattr(generation, "retrieve", lambda *a, **k: [])
    monkeypatch.setattr(generation, "_salvar_propostas_disc", lambda *a, **k: None)
    monkeypatch.setattr(
        generation.modelos, "detectar_tipo_ato", lambda t: "instrução normativa"
    )
    monkeypatch.setattr(
        generation.modelos,
        "crear_perfil_de_arquivo",
        lambda destino, texto, stem: _perfil_teste(tmp_path, conteudo),
    )

    apontamentos = [
        {"id": "ap-ort", "origem": "aprofundamento", "texto": "Corrigir erro "
         "de ortografia: 'parágrafo único do 20' deve ser 'parágrafo único do "
         "art. 20'."},
        {"id": "ap-ementa", "origem": "aprofundamento", "texto": "A ementa não "
         "menciona o fluxo financeiro do trabalho artesanal."},
    ]
    patch = {
        "numero": "INSTRUÇÃO NORMATIVA Nº 1/2026",
        "ementa": "Dispõe sobre o trabalho artesanal.",
        "alteracoes": [],
        "remocoes": [],
        "adicoes_estruturais": [],
        "apontamentos_analise": [],
    }
    monkeypatch.setattr(
        di, "_extrair_json_com_retry", lambda *a, **k: dict(patch)
    )

    resultado = json.loads(
        generation.melhorar_documento_usuario(
            filename="USUARIO_ARTESANATO.docx",
            apontamentos=apontamentos,
        )
    )
    assert resultado["status"] == "improved"

    ultima = generation.ultima_comparacao()
    passagens_analise = ultima["passagens_analise"]
    assert [p["passo"] for p in passagens_analise] == _PASSOS_ANALISE
    assert all(p["tempo_ms"] is not None for p in passagens_analise)

    cobertura = {c["apontamento_id"]: c for c in ultima["apontamentos_analise"]}
    # O achado de redação foi resolvido por correção automática.
    assert "ap-ort" in cobertura
    assert cobertura["ap-ort"]["status"] == "aplicado"
    assert "correção automática" in cobertura["ap-ort"]["motivo"]
    # A ementa ('não menciona fluxo...') não contradiz o original → acionável.
    assert "ap-ementa" in cobertura

    # Nenhuma passagem pós-geração falhou (senão cairia no fallback e o status
    # de ap-ort não seria 'aplicado').
    assert all(p.get("ok") for p in (ultima["passagens"] or []))

    # O DOCX final contém a correção inline (impossível com fallback).
    doc_final = Document(str(resultado["output_path"]))
    alvo = next(
        p for p in doc_final.paragraphs
        if "parágrafo único do 20" in p.text and "art. 20" in p.text
    )
    correcoes = [
        "".join(t.text or "" for t in r.iter(_NS + "t"))
        for r in alvo._element.iter(_NS + "r")
    ]
    assert any(t == "parágrafo único do 20" for t in correcoes)
    assert any(t == "parágrafo único do art. 20" for t in correcoes)