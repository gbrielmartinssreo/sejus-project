from fastapi.testclient import TestClient

from sejus_project.agent import agent
from sejus_project.tools import document_generation as generation
from sejus_project.web.render_html import minuta_para_html, minuta_para_texto
from sejus_project.web.server import app

ESTRUTURA = {
    "numero": "PORTARIA Nº 001/2026/GAB-SEJUS/MT",
    "ementa": "Dispõe sobre rotina de limpeza.",
    "preambulo": "O SECRETÁRIO DE ESTADO DE JUSTIÇA, no uso das atribuições,",
    "considerandos": ["CONSIDERANDO a necessidade de padronizar a rotina;"],
    "resolutivo": "RESOLVE:",
    "corpo": [
        {
            "rotulo": "Art. 1º",
            "texto": "Instituir o Programa de Limpeza.",
            "subitens": [
                {"tipo": "inciso", "rotulo": "I -", "texto": "padronizar os procedimentos;"},
            ],
        }
    ],
    "fechamento": [{"rotulo": "", "texto": "Esta Portaria entra em vigor na data de sua publicação."}],
    "local_data": "Cuiabá-MT, 8 de setembro de 2026.",
    "assinaturas": [{"nome": "VITOR HUGO BRUZULATO TEIXEIRA", "cargo": "Secretário de Estado de Justiça"}],
}


def test_minuta_para_html_escapa_texto():
    estrutura = {
        "numero": 'PORTARIA <Nº> 1 & 2 "x"',
        "ementa": "Dispõe sobre & <tag>.",
        "corpo": [
            {
                "rotulo": "Art. 1º",
                "texto": "Item com &, < e >.",
                "subitens": [{"tipo": "inciso", "rotulo": "I -", "texto": "A < B"}],
            }
        ],
        "fechamento": [{"rotulo": "", "texto": "Entra em vigor."}],
        "local_data": "Cuiabá-MT.",
        "assinaturas": [{"nome": "FULANO", "cargo": "Secretário"}],
    }

    html = minuta_para_html(estrutura)

    assert "&lt;" in html and "&amp;" in html and "&quot;" in html
    assert "<script>" not in html
    assert 'class="minuta-documento"' in html
    assert 'class="minuta-numero"' in html
    assert "Art. 1º" in html
    assert "I -" in html


def test_minuta_para_texto():
    texto = minuta_para_texto(ESTRUTURA)

    assert "PORTARIA Nº 001/2026/GAB-SEJUS/MT" in texto
    assert "Art. 1º" in texto
    assert "Cuiabá-MT, 8 de setembro de 2026." in texto


def test_index_entrega_pagina():
    client = TestClient(app)
    resp = client.get("/")

    assert resp.status_code == 200
    assert "SEJUS" in resp.text
    assert "/static/app.js" in resp.text


def test_chat_devolve_reply_e_minuta(monkeypatch, tmp_path):
    from docx import Document

    output = tmp_path / "USUARIO_Portaria_abc.docx"
    document = Document()
    document.add_paragraph("PORTARIA Nº 45/2026/GAB-SEJUS/MT")
    document.save(str(output))

    monkeypatch.setattr(agent, "executar", lambda msg: "Resposta de teste.")
    monkeypatch.setattr(generation, "consumir_geracao_do_turno", lambda: True)
    monkeypatch.setattr(
        generation,
        "ultima_minuta",
        lambda: {"estructura": ESTRUTURA, "modelo": "PORTARIA", "output_path": str(output)},
    )
    monkeypatch.setattr(generation, "has_pending_document", lambda: False)
    monkeypatch.setattr(generation, "modelo_usuario_ativo", lambda: None)

    client = TestClient(app)
    resp = client.post("/api/chat", json={"message": "gere uma portaria"})

    assert resp.status_code == 200
    dados = resp.json()
    assert dados["reply"] == "Resposta de teste."
    assert dados["minuta_nome"] == "USUARIO_Portaria_abc.docx"
    assert dados["minuta_docx"] == "/api/minuta/docx"
    assert dados["minuta_texto"]
    assert "Dispõe sobre rotina de limpeza." in dados["minuta_texto"]
    assert dados["pendente"] is False


def test_chat_repassa_modelo_usuario_ativo(monkeypatch):
    monkeypatch.setattr(agent, "executar", lambda msg: "Resposta de teste.")
    monkeypatch.setattr(generation, "ultima_minuta", lambda: None)
    monkeypatch.setattr(generation, "has_pending_document", lambda: False)
    monkeypatch.setattr(
        generation,
        "modelo_usuario_ativo",
        lambda: {"filename": "Portaria 45.docx", "modelo": "USUARIO_Portaria_45", "tipo_ato": "portaria"},
    )

    client = TestClient(app)
    dados = client.post("/api/chat", json={"message": "gere uma portaria"}).json()

    assert dados["modelo_usuario"]["filename"] == "Portaria 45.docx"
    assert dados["modelo_usuario"]["tipo_ato"] == "portaria"


def test_chat_sem_minuta_fica_pendente(monkeypatch):
    monkeypatch.setattr(agent, "executar", lambda msg: "Deseja informar os campos?")
    monkeypatch.setattr(generation, "ultima_minuta", lambda: None)
    monkeypatch.setattr(generation, "has_pending_document", lambda: True)

    client = TestClient(app)
    dados = client.post("/api/chat", json={"message": "gere uma portaria"}).json()

    assert dados["minuta_docx"] is None
    assert dados["pendente"] is True
    assert generation.CAMPOS_BASE


def test_documento_so_aparece_no_turno_em_que_foi_gerado(monkeypatch, tmp_path):
    """O cartão de download só é anexado na mensagem em que o documento foi
    gerado; respostas seguintes não repetem o cartão."""
    from docx import Document

    sinal = {"gerou": True}

    def consumir():
        valor = sinal["gerou"]
        sinal["gerou"] = False
        return valor

    output = tmp_path / "minuta.docx"
    documento = Document()
    documento.add_paragraph("PORTARIA Nº 1/2026/GAB-SEJUS/MT")
    documento.save(str(output))

    monkeypatch.setattr(agent, "executar", lambda msg: "Resposta.")
    monkeypatch.setattr(generation, "consumir_geracao_do_turno", consumir)
    monkeypatch.setattr(
        generation,
        "ultima_minuta",
        lambda: {"estructura": ESTRUTURA, "modelo": "PORTARIA", "output_path": str(output)},
    )
    monkeypatch.setattr(generation, "has_pending_document", lambda: False)
    monkeypatch.setattr(generation, "modelo_usuario_ativo", lambda: None)

    client = TestClient(app)

    primeira = client.post("/api/chat", json={"message": "gere"}).json()
    assert primeira["minuta_docx"] == "/api/minuta/docx"
    assert primeira["minuta_pdf"] == "/api/minuta/pdf"
    assert primeira["minuta_nome"] == "minuta.docx"

    segunda = client.post("/api/chat", json={"message": "obrigado"}).json()
    assert segunda["minuta_docx"] is None
    assert segunda["minuta_pdf"] is None
    assert segunda["minuta_nome"] is None


def test_upload_salva_arquivo(monkeypatch, tmp_path):
    from sejus_project.web import server

    monkeypatch.setattr(server, "IMPORTACOES_DIR", tmp_path)
    client = TestClient(app)

    resp = client.post(
        "/api/upload",
        files={"arquivo": ("minuta.txt", b"conteudo do arquivo", "text/plain")},
    )

    assert resp.status_code == 200
    assert resp.json()["filename"] == "minuta.txt"
    assert (tmp_path / "minuta.txt").read_text() == "conteudo do arquivo"


def test_upload_rejeita_formato_invalido(monkeypatch, tmp_path):
    from sejus_project.web import server

    monkeypatch.setattr(server, "IMPORTACOES_DIR", tmp_path)
    client = TestClient(app)

    resp = client.post(
        "/api/upload",
        files={"arquivo": ("malicioso.sh", b"echo oi", "text/plain")},
    )

    assert resp.status_code == 400


def test_upload_e_analise_usam_a_mesma_pasta(monkeypatch, tmp_path):
    """Upload (/api/upload) e análise (tool do agente) devem enxergar a
    mesma pasta de importacoes_usuario, mesmo mudando o diretório de CWD."""
    from sejus_project.tools import user_files
    from sejus_project.web import server

    assert server.IMPORTACOES_DIR == user_files.IMPORTACOES_DIR
    assert server.IMPORTACOES_DIR.is_absolute()

    monkeypatch.setattr(server, "IMPORTACOES_DIR", tmp_path)
    monkeypatch.setattr(user_files, "IMPORTACOES_DIR", tmp_path)
    client = TestClient(app)

    resp = client.post(
        "/api/upload",
        files={"arquivo": ("minuta_teste.txt", b"conteudo para analise", "text/plain")},
    )
    assert resp.status_code == 200

    import json

    resultado = json.loads(user_files.analisar_arquivo_usuario("minuta_teste.txt"))
    assert resultado.get("filename") == "minuta_teste.txt"
    assert "conteudo para analise" in resultado["text"]


def test_chat_sem_soffice_nao_gera_pdf(monkeypatch, tmp_path):
    """Sem LibreOffice, o cartão mostra o DOCX mas não oferece PDF."""
    from docx import Document

    from sejus_project.web import server

    document = Document()
    document.add_paragraph("DiárioOficial")
    titulo = document.add_paragraph()
    titulo.add_run("PORTARIA Nº 45/2026/GAB-SEJUS/MT")
    document.add_paragraph("GOVERNO DO ESTADO DE MATO GROSSO - SEPLAG - IOMAT")
    output = tmp_path / "minuta.docx"
    document.save(str(output))

    monkeypatch.setattr(server, "soffice_disponivel", lambda: False)
    monkeypatch.setattr(agent, "executar", lambda msg: "Resposta de teste.")
    monkeypatch.setattr(generation, "consumir_geracao_do_turno", lambda: True)
    monkeypatch.setattr(
        generation,
        "ultima_minuta",
        lambda: {"estructura": ESTRUTURA, "modelo": "PORTARIA", "output_path": str(output)},
    )
    monkeypatch.setattr(generation, "has_pending_document", lambda: False)
    monkeypatch.setattr(generation, "modelo_usuario_ativo", lambda: None)

    client = TestClient(app)
    dados = client.post("/api/chat", json={"message": "gere uma portaria"}).json()

    assert dados["minuta_nome"] == "minuta.docx"
    assert dados["minuta_docx"] == "/api/minuta/docx"
    assert dados["minuta_pdf"] is None


def test_chat_com_libreoffice_disponibiliza_pdf(monkeypatch, tmp_path):
    """Com LibreOffice, o cartão também oferece o PDF renderizado do DOCX."""
    from docx import Document

    from sejus_project.web import server

    document = Document()
    document.add_paragraph("DiárioOficial")
    output = tmp_path / "minuta.docx"
    document.save(str(output))

    monkeypatch.setattr(server, "soffice_disponivel", lambda: True)
    monkeypatch.setattr(agent, "executar", lambda msg: "Resposta de teste.")
    monkeypatch.setattr(generation, "consumir_geracao_do_turno", lambda: True)
    monkeypatch.setattr(
        generation,
        "ultima_minuta",
        lambda: {"estructura": ESTRUTURA, "modelo": "PORTARIA", "output_path": str(output)},
    )
    monkeypatch.setattr(generation, "has_pending_document", lambda: False)
    monkeypatch.setattr(generation, "modelo_usuario_ativo", lambda: None)

    client = TestClient(app)
    dados = client.post("/api/chat", json={"message": "gere uma portaria"}).json()

    assert dados["minuta_pdf"] == "/api/minuta/pdf"
    assert dados["minuta_docx"] == "/api/minuta/docx"
    assert dados["minuta_nome"] == "minuta.docx"
    assert dados["minuta_texto"]


def test_endpoints_minuta_pdf_e_docx(monkeypatch, tmp_path):
    from docx import Document

    from sejus_project.web import server

    document = Document()
    document.add_paragraph("DiárioOficial")
    docx = tmp_path / "minuta.docx"
    document.save(str(docx))
    pdf = tmp_path / "minuta.pdf"
    pdf.write_bytes(b"%PDF-1.4\nminuta renderizada")

    monkeypatch.setattr(
        generation,
        "ultima_minuta",
        lambda: {"estructura": ESTRUTURA, "modelo": "PORTARIA", "output_path": str(docx)},
    )
    monkeypatch.setattr(server, "docx_para_pdf", lambda docx_path: pdf)

    client = TestClient(app)

    resp_pdf = client.get("/api/minuta/pdf")
    assert resp_pdf.status_code == 200
    assert resp_pdf.headers["content-type"].startswith("application/pdf")
    assert resp_pdf.content[:5] == b"%PDF-"

    resp_docx = client.get("/api/minuta/docx")
    assert resp_docx.status_code == 200
    assert resp_docx.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    assert resp_docx.content == docx.read_bytes()


def test_endpoint_minuta_sem_minuta(monkeypatch):
    monkeypatch.setattr(generation, "ultima_minuta", lambda: None)
    client = TestClient(app)
    assert client.get("/api/minuta/pdf").status_code == 404
    assert client.get("/api/minuta/docx").status_code == 404


def test_chat_mensagem_vazia():
    client = TestClient(app)
    resp = client.post("/api/chat", json={"message": "   "})

    assert resp.status_code == 400
    dados = resp.json()
    assert "reply" in dados


def test_chat_mensagem_muito_longa():
    from sejus_project.web import server

    client = TestClient(app)
    resp = client.post(
        "/api/chat", json={"message": "a" * (server.MAX_MESSAGE_CHARS + 1)}
    )

    assert resp.status_code == 400
    assert "longa" in resp.json()["reply"]


def test_chat_excecao_nao_vaza_html(monkeypatch):
    def estourar(_msg):
        raise RuntimeError("Qdrant fora do ar")

    monkeypatch.setattr(agent, "executar", estourar)

    client = TestClient(app, raise_server_exceptions=False)
    resp = client.post("/api/chat", json={"message": "gere uma portaria"})

    assert resp.status_code == 500
    assert resp.headers["content-type"].startswith("application/json")
    assert "<html" not in resp.text
    dados = resp.json()
    assert "detail" in dados
    assert "Qdrant fora do ar" in dados["error"]


def test_limpar_conversa_reseta_estado(monkeypatch):
    monkeypatch.setattr(agent, "limpar_conversa", lambda: None)
    client = TestClient(app)
    resp = client.post("/api/conversa/limpar")

    assert resp.status_code == 200
    assert resp.json()["detail"] == "Conversa limpa."


def _docx_portaria(bytes_conteudo=None) -> bytes:
    import io

    from docx import Document

    document = Document()
    titulo = document.add_paragraph()
    titulo.add_run("PORTARIA Nº 45/2026/GAB-SEJUS/MT")
    document.add_paragraph("Art. 1º Instituir o Programa de Limpeza.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_enviar_modelo_define_estado(monkeypatch, tmp_path):
    from sejus_project.web import server

    monkeypatch.setattr(server, "IMPORTACOES_DIR", tmp_path)
    import sejus_project.tools.document_generation as generation_mod

    generation_mod._modelo_usuario = None
    client = TestClient(app)

    resp = client.post(
        "/api/modelo",
        files={"arquivo": ("Portaria Modelo.docx", _docx_portaria(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert resp.status_code == 200
    dados = resp.json()
    assert dados["filename"] == "Portaria Modelo.docx"
    assert dados["tipo_ato"] == "portaria"
    assert (tmp_path / "Portaria Modelo.docx").is_file()
    assert generation_mod._modelo_usuario is not None
    assert generation_mod._modelo_usuario["perfil"].name == "USUARIO_Portaria_Modelo"
    generation_mod._modelo_usuario = None


def test_enviar_modelo_rejeita_pdf(monkeypatch, tmp_path):
    from sejus_project.web import server

    monkeypatch.setattr(server, "IMPORTACOES_DIR", tmp_path)
    client = TestClient(app)

    resp = client.post(
        "/api/modelo",
        files={"arquivo": ("ato.pdf", b"%PDF-1.4 falso", "application/pdf")},
    )

    assert resp.status_code == 400
    assert "docx" in resp.json()["detail"].casefold()