from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from sejus_project.tools import document_generation as generation
from sejus_project.tools import minuta, modelos
from sejus_project.tools.render_docx import docx_para_html

ESTRUTURA = {
    "numero": "PORTARIA Nº 001/2026/GAB-SEJUS/MT",
    "ementa": "Dispõe sobre rotina de limpeza.",
    "preambulo": "O SECRETÁRIO DE ESTADO DE JUSTIÇA, no uso das atribuições,",
    "considerandos": ["CONSIDERANDO a necessidade de padronizar a rotina;"],
    "resolutivo": "RESOLVE:",
    "corpo": [
        {
            "rotulo": "Art. 1º",
            "texto": "Instituir o Programa de Limpeza nas unidades da SEJUS.",
            "subitens": [
                {"tipo": "inciso", "rotulo": "I -", "texto": "padronizar os procedimentos diários;"},
            ],
        }
    ],
    "fechamento": [{"rotulo": "Art. 2º", "texto": "Esta Portaria entra em vigor na data de sua publicação."}],
    "local_data": "Cuiabá-MT, 8 de setembro de 2026.",
    "assinaturas": [{"nome": "VITOR HUGO BRUZULATO TEIXEIRA", "cargo": "Secretário de Estado de Justiça"}],
}


def test_docx_para_html_preserva_formatacao(tmp_path):
    document = Document()
    document.add_paragraph("18 de junho de 2025")
    document.add_paragraph("DiárioOficial")
    document.add_paragraph("Nº 29.013")
    document.add_paragraph()
    titulo = document.add_paragraph()
    titulo.add_run("PORTARIA Nº 45/2025/GAB-SEJUS/MT")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run("Texto de exemplo.")
    run.font.bold = True
    run.font.size = Pt(11)
    cab_titulo = document.add_paragraph(style="Heading 2")
    cab_titulo.add_run("TÍTULO COM ESTILO DE TÍTULO")
    estilo_titulo = document.styles["Heading 2"]
    estilo_titulo.font.name = "Arial"
    estilo_titulo.font.size = Pt(12)
    document.add_paragraph(
        "GOVERNO DO ESTADO DE MATO GROSSO Secretaria de Estado de Planejamento "
        "e Gestão - SEPLAG - Imprensa Oficial - IOMAT"
    )
    caminho = tmp_path / "documento.docx"
    document.save(str(caminho))

    html = docx_para_html(caminho)

    assert html.startswith('<div class="minuta-documento"')
    assert "DiárioOficial" in html
    assert ">&nbsp;</p>" in html
    assert "text-align:justify" in html
    assert "font-weight:700" in html
    assert "font-size:11.0pt" in html
    assert "GOVERNO DO ESTADO DE MATO GROSSO" in html
    assert "<script" not in html

    # Herança de estilos: o parágrafo usa Heading 2 (Arial, negrito, 12pt)
    # mesmo sem propriedades diretas nos runs.
    idx_titulo = html.index("TÍTULO COM ESTILO DE TÍTULO")
    span_titulo = html[html.rfind("<span", 0, idx_titulo) : idx_titulo]
    assert "font-size:12.0pt" in span_titulo
    assert "font-weight:700" in span_titulo
    assert "Arial" in span_titulo
    # Props de parágrafo não vazam para dentro do span.
    assert "text-align:" not in span_titulo
    assert "margin-" not in span_titulo


def test_preview_end_to_end_igual_output(tmp_path):
    """O HTML do front, gerado a partir do DOCX, contém cabeçalho, corpo novo
    e rodapé preservados pelo modelo do usuário."""
    document = Document()
    document.add_paragraph("18 de junho de 2025")
    document.add_paragraph("DiárioOficial")
    document.add_paragraph("Nº 29.013")
    titulo = document.add_paragraph()
    titulo.add_run("PORTARIA Nº 45/2025/GAB-SEJUS/MT")
    document.add_paragraph("O Secretário de Estado de Justiça, no uso das atribuições,")
    document.add_paragraph("Art. 1º Instituir o Comitê de Proteção de Dados.")
    document.add_paragraph("CONTEÚDO DE OUTRO ATO QUE DEVE SUMIR")
    document.add_paragraph(
        "GOVERNO DO ESTADO DE MATO GROSSO Secretaria de Estado de Planejamento "
        "e Gestão - SEPLAG - Imprensa Oficial - IOMAT"
    )
    caminho = tmp_path / "Modelo_Diario.docx"
    document.save(str(caminho))

    texto = generation.extract_file_text(caminho)
    perfil = modelos.crear_perfil_de_arquivo(caminho, texto, "Modelo_Diario")
    output_path = minuta.montar_docx(perfil, ESTRUTURA, tmp_path)

    html = docx_para_html(output_path)

    assert "DiárioOficial" in html
    assert "Nº 29.013" in html
    assert "PORTARIA Nº 001/2026/GAB-SEJUS/MT" in html
    assert "Programa de Limpeza nas unidades da SEJUS." in html
    assert "VITOR HUGO BRUZULATO TEIXEIRA" in html
    assert "GOVERNO DO ESTADO DE MATO GROSSO" in html
    assert "CONTEÚDO DE OUTRO ATO" not in html
    assert "Comitê de Proteção de Dados" not in html