import pytest
from docx import Document

from sejus_project.tools import pdf_preview

pytestmark = pytest.mark.skipif(
    not pdf_preview.soffice_disponivel(),
    reason="LibreOffice não está disponível nesta máquina.",
)


def test_docx_para_pdf_gera_pdf(tmp_path):
    docx = tmp_path / "minuta.docx"
    document = Document()
    document.add_paragraph("DiárioOficial")
    document.add_paragraph("PORTARIA Nº 45/2026/GAB-SEJUS/MT")
    document.save(str(docx))

    pdf = pdf_preview.docx_para_pdf(docx)

    assert pdf is not None
    assert pdf.is_file()
    assert pdf.read_bytes()[:5] == b"%PDF-"


def test_docx_para_pdf_cacheia(tmp_path):
    docx = tmp_path / "minuta.docx"
    document = Document()
    document.add_paragraph("Conteúdo de teste.")
    document.save(str(docx))

    primeiro = pdf_preview.docx_para_pdf(docx)
    segundo = pdf_preview.docx_para_pdf(docx)

    assert primeiro == segundo
    assert primeiro.exists()